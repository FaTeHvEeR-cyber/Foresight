"""Pytest fixtures and test data generators for Foresight backend tests."""

import io
import pytest
import pandas as pd
import numpy as np
from fastapi.testclient import TestClient
from pypdf import PdfWriter
import docx
import openpyxl

from app.main import app
from app.config import settings


@pytest.fixture
def client():
    """FastAPI TestClient fixture."""
    return TestClient(app)


@pytest.fixture
def sample_tabular_df():
    """Sample DataFrame matching ingestion, profiling, and modeling tests."""
    return pd.DataFrame({
        "age": [25.0, 30.0, np.nan, 40.0, 50.0],
        "salary": [50000.0, 60000.0, 70000.0, np.nan, 90000.0],
        "department": ["Eng", "Sales", "HR", "Eng", None],
        "is_active": [True, False, True, True, False],
    })


@pytest.fixture
def sample_csv_bytes(sample_tabular_df):
    """CSV bytes generated from sample DataFrame (5 rows, 4 columns)."""
    buf = io.StringIO()
    sample_tabular_df.to_csv(buf, index=False)
    return buf.getvalue().encode("utf-8")


@pytest.fixture
def sample_tsv_bytes(sample_tabular_df):
    """TSV bytes generated from sample DataFrame."""
    buf = io.StringIO()
    sample_tabular_df.to_csv(buf, sep="\t", index=False)
    return buf.getvalue().encode("utf-8")


@pytest.fixture
def sample_xlsx_bytes(sample_tabular_df):
    """Valid XLSX bytes generated via openpyxl."""
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        sample_tabular_df.to_excel(writer, index=False)
    return buf.getvalue()


@pytest.fixture
def sample_parquet_bytes():
    """Valid Parquet bytes with 4 rows and 3 columns."""
    df = pd.DataFrame({
        "col_a": [1.0, 2.0, 3.0, 4.0],
        "col_b": [10, 20, 30, 40],
        "col_c": ["x", "y", "z", "w"],
    })
    buf = io.BytesIO()
    df.to_parquet(buf, index=False)
    return buf.getvalue()


@pytest.fixture
def sample_pdf_bytes():
    """Valid PDF bytes generated cleanly via pypdf."""
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=144)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


@pytest.fixture
def sample_docx_bytes():
    """Valid DOCX bytes generated via python-docx."""
    doc = docx.Document()
    doc.add_heading("Foresight Research Dossier", level=1)
    doc.add_paragraph("This is a sample document for testing Engine C ingestion.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
def sample_txt_bytes():
    """Valid plain text bytes."""
    return b"Foresight plain text report content.\nLine 2: local intelligence."


@pytest.fixture
def sample_md_bytes():
    """Valid markdown bytes."""
    return b"# Markdown Executive Report\n\n- Point A\n- Point B\n"


# ---------------------------------------------------------------------------
# Phase 3A Fixtures
# ---------------------------------------------------------------------------

@pytest.fixture(scope="session")
def bike_df():
    rng = np.random.default_rng(0)
    d = pd.date_range("2018-01-01", "2019-12-31", freq="D")
    t = np.arange(len(d))
    reg = 2000 + 3 * t + 500 * np.sin(2 * np.pi * t / 365) + 300 * (d.dayofweek < 5) + rng.normal(0, 120, len(d))
    cas = 500 + 200 * np.sin(2 * np.pi * t / 365) + 250 * (d.dayofweek >= 5) + rng.normal(0, 60, len(d))
    return pd.DataFrame({
        "instant": t + 1,
        "dteday": d.strftime("%d-%m-%Y"),
        "season": (d.month % 12) // 3 + 1,
        "yr": (d.year - 2018),
        "mnth": d.month,
        "holiday": (np.arange(len(d)) % 50 == 0).astype(int),
        "weekday": d.dayofweek,
        "temp": 0.5 + 0.3 * np.sin(2 * np.pi * t / 365),
        "casual": cas.round(),
        "registered": reg.round(),
        "cnt": (cas.round() + reg.round()),
    })


@pytest.fixture(scope="session")
def airline_df():
    d = pd.date_range("1949-01-01", periods=144, freq="MS")
    t = np.arange(144)
    v = (110 + 2.6 * t) * (1 + 0.25 * np.sin(2 * np.pi * t / 12))
    return pd.DataFrame({"Month": d.strftime("%Y-%m"), "#Passengers": v.round().astype(int)})


@pytest.fixture(scope="session")
def retail_df():
    rng = np.random.default_rng(1)
    days = pd.date_range("2010-12-01", "2011-12-09", freq="D")
    days = days[days.dayofweek != 5]  # no Saturday trading
    rows = []
    for i, d in enumerate(days):
        k = rng.integers(20, 60)
        for j in range(k):
            inv = f"{500000 + i * 100 + j // 5}"
            qty = int(rng.integers(1, 12))
            price = float(np.round(rng.uniform(0.5, 9.0), 2))
            rows.append((inv, "X", "Widget", qty, f"{d.month}/{d.day}/{d.year} {rng.integers(8, 17)}:{rng.integers(10, 59)}", price))
    df = pd.DataFrame(rows, columns=["InvoiceNo", "StockCode", "Description", "Quantity", "InvoiceDate", "UnitPrice"])
    cancel = df.sample(300, random_state=2).index
    df.loc[cancel, "InvoiceNo"] = "C" + df.loc[cancel, "InvoiceNo"]
    df.loc[cancel, "Quantity"] = -df.loc[cancel, "Quantity"]
    df.loc[df.sample(80, random_state=3).index, "UnitPrice"] = 0.0
    df.insert(0, "index", np.arange(len(df)))
    return df


@pytest.fixture(scope="session")
def wholesale_df():
    rng = np.random.default_rng(4)
    n = 440
    ch = rng.choice([1, 2], n, p=[0.68, 0.32])
    reg = rng.choice([1, 2, 3], n, p=[0.72, 0.18, 0.10])
    return pd.DataFrame({
        "Channel": ch,
        "Region": reg,
        "Fresh": rng.lognormal(8.5, 1.0, n) * np.where(ch == 1, 1.2, 0.7),
        "Milk": rng.lognormal(7.8, 1.0, n) * np.where(ch == 2, 2.0, 1.0),
        "Grocery": rng.lognormal(8.0, 1.1, n),
    })


@pytest.fixture(scope="session")
def promo_df():
    rng = np.random.default_rng(5)
    d = pd.date_range("2014-01-01", periods=600, freq="D")
    promo = rng.integers(0, 2, len(d))
    return pd.DataFrame({
        "Date": d.strftime("%Y-%m-%d"),
        "Sales": (5000 * (1 + 0.387 * promo) + rng.normal(0, 400, len(d))).round(),
        "Promo": promo,
    })


def to_csv_bytes(df, encoding="utf-8"):
    return df.to_csv(index=False).encode(encoding)
